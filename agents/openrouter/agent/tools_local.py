# kAIm56 — self-hosted Firecracker AI-agent platform
# Copyright (C) 2026 the kAIm56 authors
# SPDX-License-Identifier: AGPL-3.0-or-later
# This program is free software under the GNU AGPL v3+; see LICENSE.
"""Tools that act inside the VM: shell, files, Office output (xlsx/docx), directory listing, HTTP fetch, PDF reading, web search.

Part of the openrouter agent package (runs inside the VM): no import from the package root. Sibling modules are used
as ``_name.func`` (module attribute), so a test can replace one definition in
one place.
"""
import json
import os
import re
import subprocess
import urllib.error

from . import config as _config
from . import mgrclient as _mgrclient


# --- built-in tools ---------------------------------------------------------
def t_bash(command):
    p = subprocess.run(command, shell=True, cwd=_config.WORKDIR, capture_output=True,
                       text=True, timeout=_config.BASH_TIMEOUT)
    return (p.stdout + p.stderr).strip() or f"(exit {p.returncode}, no output)"


def _safe(path):
    p = os.path.abspath(os.path.join(_config.WORKDIR, path)) if not os.path.isabs(path) else path
    return p


def t_read_file(path):
    with open(_safe(path)) as f:
        return f.read(_config.MAX_TOOL_OUT)


def t_write_file(path, content):
    p = _safe(path)
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    with open(p, "w") as f:
        f.write(content)
    return f"written: {p} ({len(content)} characters)"


# --- Office output (openpyxl / python-docx, both in the rootfs) -------------
def _rows_norm(rows):
    """Rows for a sheet: a JSON string, a list of lists, or a list of dicts
    (dict keys become the header row). Cells are left as str/int/float."""
    if isinstance(rows, str):
        rows = json.loads(rows)
    if not isinstance(rows, list):
        raise ValueError("rows must be a list of rows (lists) or of objects")
    if rows and isinstance(rows[0], dict):
        keys = []
        for r in rows:
            for k in r:
                if k not in keys:
                    keys.append(k)
        return [keys] + [[r.get(k, "") for k in keys] for r in rows]
    return [list(r) if isinstance(r, (list, tuple)) else [r] for r in rows]


def t_write_xlsx(path, rows, sheet="Sheet1"):
    """Write a spreadsheet (.xlsx) into the workspace. rows: list of lists
    (first row = header) or list of objects; also accepted as a JSON string."""
    try:
        import openpyxl
    except ImportError:
        return "⚠️ write_xlsx needs openpyxl in this image (rebuild the rootfs)"
    try:
        data = _rows_norm(rows)
    except (ValueError, TypeError) as e:
        return f"⚠️ rows: {e}"
    p = _safe(path if str(path).lower().endswith(".xlsx") else f"{path}.xlsx")
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = str(sheet or "Sheet1")[:31]
    for r in data:
        ws.append([c if isinstance(c, (int, float)) or c is None else str(c) for c in r])
    if data:
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        for col in ws.columns:
            w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, w + 2))
    wb.save(p)
    return f"written: {p} ({max(0, len(data) - 1)} rows, {len(data[0]) if data else 0} columns)"


def _md_blocks(md):
    """A small Markdown subset -> blocks for a document: ('h', level, text),
    ('li', text), ('p', text). Inline **bold** stays as markers for the writer."""
    out, para = [], []
    def flush():
        if para:
            out.append(("p", " ".join(para))); para.clear()
    for line in str(md or "").splitlines():
        t = line.rstrip()
        if not t.strip():
            flush(); continue
        m = re.match(r"^(#{1,3})\s+(.*)$", t)
        if m:
            flush(); out.append(("h", len(m.group(1)), m.group(2).strip())); continue
        m = re.match(r"^\s*[-*]\s+(.*)$", t)
        if m:
            flush(); out.append(("li", m.group(1).strip())); continue
        para.append(t.strip())
    flush()
    return out


def t_write_docx(path, markdown, title=None):
    """Write a Word document (.docx) into the workspace from a small Markdown
    subset: # headings (1-3), - bullets, paragraphs, **bold** inline."""
    try:
        import docx
    except ImportError:
        return "⚠️ write_docx needs python-docx in this image (rebuild the rootfs)"
    p = _safe(path if str(path).lower().endswith(".docx") else f"{path}.docx")
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    d = docx.Document()
    if title:
        d.add_heading(str(title), 0)
    def runs(par, text):
        for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if seg:
                par.add_run(seg).bold = bool(i % 2)
    blocks = _md_blocks(markdown)
    for b in blocks:
        if b[0] == "h":
            d.add_heading(b[2], min(3, b[1]))
        elif b[0] == "li":
            runs(d.add_paragraph(style="List Bullet"), b[1])
        else:
            runs(d.add_paragraph(), b[1])
    d.save(p)
    return f"written: {p} ({len(blocks)} blocks)"


def t_list_dir(path="."):
    return "\n".join(sorted(os.listdir(_safe(path)))) or "(empty)"


def _html_to_text(html):
    """Readable text out of an HTML page: scripts/styles gone, block tags as
    line breaks, entities resolved, whitespace collapsed. Links keep their
    target in brackets so the model can follow them with another fetch."""
    import html as _h
    t = html.replace("\r", "")
    t = re.sub(r"(?is)<(script|style|noscript|svg|head)[^>]*>.*?</\1>", " ", t)
    t = re.sub(r'(?is)<a[^>]*href=["\'](https?://[^"\']+)["\'][^>]*>(.*?)</a>',
               lambda m: re.sub(r"<[^>]+>", "", m.group(2)) + " [" + m.group(1) + "]", t)
    t = re.sub(r"(?i)<(br|/p|/div|/li|/tr|/h[1-6]|/section|/article)[^>]*>", "\n", t)
    t = re.sub(r"(?s)<[^>]+>", " ", t)
    t = _h.unescape(t)
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n[ \t]*", "\n", t)
    return re.sub(r"\n{3,}", "\n\n", t).strip()


def t_http_fetch(url, method="GET", raw=False):
    """Fetch a URL. HTML is converted to readable text (a modern page is 90%
    markup and scripts — hard-truncated raw HTML used to cut content off before
    it ever appeared, and the model concluded pages were "too complex").
    raw=true returns the unconverted body for the cases that need markup."""
    req = urllib.request.Request(url, method=method, headers={
        # A browser UA: big job/news portals 403 obvious bot agents outright.
        "User-Agent": "Mozilla/5.0 (X11; Linux x86_64; rv:120.0) "
                      "Gecko/20100101 Firefox/120.0",
        "Accept-Language": "de,en;q=0.7"})
    r = urllib.request.urlopen(req, timeout=30)
    body = r.read(5_000_000).decode("utf-8", "replace")
    ctype = (r.headers.get("Content-Type") or "").lower()
    if raw or not ("html" in ctype or body.lstrip()[:200].lower().startswith(("<!doctype", "<html"))):
        return body
    return _html_to_text(body)


def t_read_pdf(path, pages=""):
    """Extract PDF text via pdftotext. `path` = workspace file OR
    http(s) URL. `pages` optional as a range, e.g. '1-5'."""
    import re
    import tempfile
    tmp = None
    try:
        if str(path).startswith(("http://", "https://")):
            data = urllib.request.urlopen(
                urllib.request.Request(path, headers={"User-Agent": "or-agent"}),
                timeout=30).read(50 * 1024 * 1024)
            tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
            tmp.write(data)
            tmp.close()
            src = tmp.name
        else:
            src = _safe(path)
        cmd = ["pdftotext", "-q"]
        m = re.match(r"(\d+)-(\d+)$", (pages or "").strip())
        if m:
            cmd += ["-f", m.group(1), "-l", m.group(2)]
        cmd += [src, "-"]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
        txt = (r.stdout or "").strip()
        if not txt:
            return ("PDF error: " + (r.stderr.strip()[:300] or "")
                    if r.returncode != 0
                    else "(no text in the PDF — possibly a scanned image without a text layer)")
        return txt[:_config.MAX_TOOL_OUT]
    except Exception as e:
        return f"Error: {e!r}"
    finally:
        if tmp:
            try:
                os.remove(tmp.name)
            except OSError:
                pass


def _ddg_search(query, count):
    """DuckDuckGo HTML. Returns a result list, or None when DDG serves its
    bot challenge instead of results (HTTP 202 + "anomaly" page — since
    2026-09 the norm for datacenter IPs, found via a user's empty search)."""
    q = urllib.parse.quote(query)
    req = urllib.request.Request(
        "https://html.duckduckgo.com/html/?q=" + q,
        headers={"User-Agent": "Mozilla/5.0 (X11; Linux x86_64; rv:120.0) Gecko/20100101 Firefox/120.0"})
    r = urllib.request.urlopen(req, timeout=30)
    html = r.read(1_000_000).decode("utf-8", "replace")
    if r.status != 200 or "anomaly" in html[:4000] or "challenge" in html[:4000]:
        return None                       # blocked -> let the caller try elsewhere
    hrefs = re.findall(r'class="result__a"[^>]*href="([^"]+)"', html)
    titles = re.findall(r'class="result__a"[^>]*>(.*?)</a>', html, re.S)
    snips = re.findall(r'class="result__snippet"[^>]*>(.*?)</a>', html, re.S)

    def real(h):
        m = re.search(r"uddg=([^&]+)", h)
        return urllib.parse.unquote(m.group(1)) if m else h

    return [(_ws_clean(titles[i]) if i < len(titles) else "",
             real(hrefs[i]),
             _ws_clean(snips[i]) if i < len(snips) else "")
            for i in range(min(count, len(hrefs)))]


def _bing_search(query, count):
    """Bing HTML. Result URLs are /ck/a redirects carrying the target
    base64-encoded in u=a1<payload> — decoded here."""
    q = urllib.parse.quote(query)
    req = urllib.request.Request(
        "https://www.bing.com/search?q=" + q + "&count=" + str(max(count, 10)),
        headers={"User-Agent": "Mozilla/5.0 (X11; Linux x86_64; rv:120.0) Gecko/20100101 Firefox/120.0",
                 "Accept-Language": "de,en;q=0.7"})
    html = urllib.request.urlopen(req, timeout=30).read(2_000_000).decode("utf-8", "replace")
    out = []
    # Per result block, not one regex across the page: the snippet <p> sits at
    # varying depths and a greedy pattern either misses it or bleeds across
    # blocks.
    for block in html.split('<li class="b_algo"')[1:]:
        block = block.split("</li>", 1)[0]
        a = re.search(r'<h2[^>]*>.*?<a[^>]*href="([^"]+)"[^>]*>(.*?)</a>', block, re.S)
        if not a:
            continue
        p = re.search(r"<p[^>]*>(.*?)</p>", block, re.S)
        out.append((_ws_clean(a.group(2)), _bing_real_url(a.group(1)),
                    _ws_clean(p.group(1)) if p else ""))
        if len(out) >= count:
            break
    return out


def _bing_real_url(href):
    """Bing /ck/a redirect -> target URL (u=a1<urlsafe-base64>)."""
    import base64
    h = urllib.parse.unquote(href.replace("&amp;", "&"))
    m = re.search(r"[&?]u=a1([A-Za-z0-9_\-]+)", h)
    if not m:
        return h
    raw = m.group(1)
    try:
        return base64.urlsafe_b64decode(raw + "=" * (-len(raw) % 4)).decode(
            "utf-8", "replace")
    except Exception:
        return h


def _ws_clean(s):
    return re.sub(r"<[^>]+>", "", s).replace("&amp;", "&").replace("&#x27;", "'").strip()


def t_web_search(query, count=5):
    """Web search. First choice: the manager's /api/websearch — it holds the
    Brave API key (never enters the VM) and falls back to DDG/Bing itself.
    The direct backends below only run when the manager route is missing
    (older manager). "no results" means the query found nothing; a dead
    backend must SAY so — otherwise the model concludes the thing searched
    for does not exist."""
    try:
        count = int(count)
    except (TypeError, ValueError):
        count = 5
    try:
        q = urllib.parse.urlencode({"q": query, "count": count})
        d = json.loads(_mgrclient._mgr_get(_mgrclient._manager_base(), "/api/websearch?" + q))
        if d.get("result"):
            return d["result"]
        if d.get("error"):
            raise ValueError(d["error"])
    except Exception:
        pass                      # manager route missing/broken -> go direct
    errors = []
    for name, backend in (("duckduckgo", _ddg_search), ("bing", _bing_search)):
        try:
            rows = backend(query, count)
        except Exception as e:
            errors.append(f"{name}: {e!r}")
            continue
        if rows is None:
            errors.append(f"{name}: blocked (bot challenge)")
            continue
        if rows:
            return "\n".join(f"{i+1}. {t}\n   {u}\n   {sn}"
                              for i, (t, u, sn) in enumerate(rows))
        return "no results"
    return ("⚠️ web search unavailable — every backend failed ("
            + "; ".join(errors) + "). This is an infrastructure problem, "
            "NOT an empty result: tell the user instead of concluding "
            "nothing exists.")
